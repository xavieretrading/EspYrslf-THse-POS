"""
Enhanced Generator for Espresso Yourself & Tea House Inventory Audit Template.
Generates a multi-tab, beautifully styled, color-coded, formula-powered Excel workbook
for Cebu Branch (Branch 30).
"""

import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

# Output paths
rec_dir = r"c:\Users\Philippines Freight\MainSystems\POS\recordsExcel"
workspace_dir = r"c:\Users\Philippines Freight\MainSystems\POS"
os.makedirs(rec_dir, exist_ok=True)
excel_output_path = os.path.join(rec_dir, "Espresso_Branch_Physical_Inventory_Template.xlsx")
workspace_copy_path = os.path.join(workspace_dir, "Espresso_Branch_Inventory_Template.xlsx")

# --- MASTER CATEGORIES & ITEM DATA ---
CATEGORIES_DATA = [
    {
        "cat_id": 1,
        "cat_name": "1. Coffee Beans & Drink Bases",
        "icon": "☕",
        "header_fill": "3D2314", # Rich Dark Espresso
        "items": [
            {"sku": "ING-001", "name": "Concept Blend 1 (Espresso Beans)", "unit": "Pack (1kg)", "cost": 980.00, "stock": 6.27, "notes": "Main house espresso blend"},
            {"sku": "ING-002", "name": "Australia's Own Coffee Blend", "unit": "Pack", "cost": 204.00, "stock": 1.00, "notes": "Specialty roast"},
            {"sku": "ING-003", "name": "Easy Dark Chocolate Base 1kg", "unit": "Pack (1kg)", "cost": 396.00, "stock": 1.00, "notes": "Mocha & Hot Choco base powder"},
            {"sku": "ING-004", "name": "Easy Cookies & Cream Powder 1kg", "unit": "Pack (1kg)", "cost": 305.00, "stock": 1.00, "notes": "Blended frappe & Oreo base"},
            {"sku": "ING-005", "name": "Bakersfield Caramel Fudge 1kg", "unit": "Pack (1kg)", "cost": 415.00, "stock": 1.00, "notes": "Caramel drizzle & drink base"},
            {"sku": "ING-006", "name": "Whoop Classic Ube Taro Powder 1kg", "unit": "Pack (1kg)", "cost": 380.00, "stock": 1.00, "notes": "Specialty blended drink base"},
        ]
    },
    {
        "cat_id": 2,
        "cat_name": "2. Milks & Dairy Supplies",
        "icon": "🥛",
        "header_fill": "5C3A21", # Latte Brown
        "items": [
            {"sku": "ING-007", "name": "Arla Full Cream Milk 1L", "unit": "Pack (1L)", "cost": 99.00, "stock": 10.00, "notes": "Primary steaming milk for lattes"},
            {"sku": "ING-008", "name": "Master Barista Cow's Milk 1L", "unit": "Pack (1L)", "cost": 110.00, "stock": 12.00, "notes": "Barista specialty milk"},
            {"sku": "ING-009", "name": "Doreen Condensed Milk 390g", "unit": "Can (390g)", "cost": 44.00, "stock": 19.23, "notes": "Spanish latte sweet base"},
            {"sku": "ING-010", "name": "Soymilk 1L (Plant-based)", "unit": "Pack (1L)", "cost": 87.50, "stock": 2.00, "notes": "Non-dairy milk substitute"},
            {"sku": "ING-011", "name": "Whipping Cream 1L", "unit": "Bot (1L)", "cost": 250.00, "stock": 1.00, "notes": "Cold foam & sea salt cream topper"},
            {"sku": "ING-012", "name": "Artisanal Salted Cream Cheese", "unit": "Pack", "cost": 200.00, "stock": 2.00, "notes": "Sea salt cheese foam cap"},
        ]
    },
    {
        "cat_id": 3,
        "cat_name": "3. Flavor Syrups, Purées & Sauces",
        "icon": "🍯",
        "header_fill": "7B4B28", # Warm Caramel Amber
        "items": [
            {"sku": "SYR-001", "name": "SHOTT Vanilla Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 0.75, "notes": "Vanilla Latte & Frappe"},
            {"sku": "SYR-002", "name": "SHOTT Caramel Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 0.20, "notes": "Caramel Macchiato"},
            {"sku": "SYR-003", "name": "SHOTT Macadamia Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 1.00, "notes": "Macadamia Latte"},
            {"sku": "SYR-004", "name": "SHOTT Irish Cream Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 1.00, "notes": "Irish Cream Latte"},
            {"sku": "SYR-005", "name": "SHOTT Butterscotch Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 1.00, "notes": "Butterscotch Latte"},
            {"sku": "SYR-006", "name": "SHOTT White Chocolate Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 9.84, "notes": "White Choco Latte & Matcha"},
            {"sku": "SYR-007", "name": "SHOTT Pink Guava Syrup 1L", "unit": "Bot (1L)", "cost": 890.00, "stock": 1.00, "notes": "Guava Fruit Tea & Soda"},
            {"sku": "SYR-008", "name": "SHOTT Mango Syrup 1L", "unit": "Bot (1L)", "cost": 890.00, "stock": 0.25, "notes": "Mango Tea & Soda"},
            {"sku": "SYR-009", "name": "SHOTT Pomegranate Syrup 1L", "unit": "Bot (1L)", "cost": 890.00, "stock": 0.37, "notes": "Pomegranate Tea & Soda"},
            {"sku": "SYR-010", "name": "SHOTT Lychee Syrup 1L", "unit": "Bot (1L)", "cost": 690.00, "stock": 0.83, "notes": "Lychee Tea & Soda"},
            {"sku": "SYR-011", "name": "Torani Salted Caramel Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Secondary caramel syrup"},
            {"sku": "SYR-012", "name": "Torani Macadamia Nut Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Nutty espresso drinks"},
            {"sku": "SYR-013", "name": "Torani Lychee Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Fruit soda / iced tea"},
            {"sku": "SYR-014", "name": "Torani Pomegranate Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Fruit soda / iced tea"},
            {"sku": "SYR-015", "name": "Torani Mango Puree Blend 1.89L", "unit": "Gal (1.89L)", "cost": 950.00, "stock": 1.00, "notes": "Specialty puree & smoothies"},
            {"sku": "SYR-016", "name": "Torani Strawberry Puree Blend 1.89L", "unit": "Gal (1.89L)", "cost": 950.00, "stock": 1.00, "notes": "Specialty puree & smoothies"},
            {"sku": "SYR-017", "name": "French Vanilla Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "French Vanilla Coffee"},
            {"sku": "SYR-018", "name": "Hazelnut Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Hazelnut Latte"},
            {"sku": "SYR-019", "name": "Irish Cream Syrup 750ml", "unit": "Bot (750ml)", "cost": 550.00, "stock": 1.00, "notes": "Specialty flavor"},
        ]
    },
    {
        "cat_id": 4,
        "cat_name": "4. Matcha & Specialty Teas",
        "icon": "🍵",
        "header_fill": "2D5A27", # Matcha Forest Green
        "items": [
            {"sku": "TEA-001", "name": "Matcha Powder 1kg (Culinary Grade)", "unit": "Pack (1kg)", "cost": 1200.00, "stock": 0.87, "notes": "Matcha Latte & Seasalt"},
            {"sku": "TEA-002", "name": "Easy Matcha Pure Powder 100g", "unit": "Pack (100g)", "cost": 210.00, "stock": 3.00, "notes": "Ceremonial / pure grade"},
            {"sku": "TEA-003", "name": "Possmei Jasmine Green Tea Leaves 600g", "unit": "Pack (600g)", "cost": 768.00, "stock": 1.00, "notes": "Premium tea brewing"},
            {"sku": "TEA-004", "name": "Possmei Assam Black Tea Leaves 600g", "unit": "Pack (600g)", "cost": 440.00, "stock": 1.00, "notes": "Premium black tea brewing"},
            {"sku": "TEA-005", "name": "Jasmine Tea Leaves 1kg", "unit": "Pack (1kg)", "cost": 400.00, "stock": 0.92, "notes": "Bulk brewed jasmine tea"},
            {"sku": "TEA-006", "name": "Black Tea Leaves 1kg", "unit": "Pack (1kg)", "cost": 400.00, "stock": 0.98, "notes": "Bulk brewed black tea"},
            {"sku": "TEA-007", "name": "Dilmah Earl Grey Tea Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 91.00, "notes": "Hot brewed specialty tea"},
            {"sku": "TEA-008", "name": "Dilmah Gourmet Green Tea Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 100.00, "notes": "Hot brewed specialty tea"},
            {"sku": "TEA-009", "name": "Dilmah Green Tea with Jasmine Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 100.00, "notes": "Hot brewed specialty tea"},
            {"sku": "TEA-010", "name": "Dilmah Pure Camomile Flowers Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 100.00, "notes": "Hot brewed herbal tea"},
            {"sku": "TEA-011", "name": "Dilmah Pure Peppermint Leaves Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 98.00, "notes": "Hot brewed herbal tea"},
            {"sku": "TEA-012", "name": "Dilmah Traditional Oolong Tea Bags", "unit": "Bag / Sachet", "cost": 45.00, "stock": 100.00, "notes": "Hot brewed specialty tea"},
        ]
    },
    {
        "cat_id": 5,
        "cat_name": "5. Sweeteners, Toppings & Mixers",
        "icon": "🍬",
        "header_fill": "996515", # Warm Toffee Gold
        "items": [
            {"sku": "SWT-001", "name": "Hi-Fructose Corn Syrup 500ml", "unit": "Bot (500ml)", "cost": 94.00, "stock": 10.00, "notes": "Primary liquid sweetener"},
            {"sku": "SWT-002", "name": "Refined White Sugar 1kg", "unit": "Pack (1kg)", "cost": 80.00, "stock": 1.00, "notes": "Syrup prep & table sugar"},
            {"sku": "SWT-003", "name": "Equal Sweetener Sachets (Box of 50s)", "unit": "Pack (50s)", "cost": 167.50, "stock": 1.00, "notes": "Low calorie option"},
            {"sku": "SWT-004", "name": "Lotus Biscoff Spread Smooth 400g", "unit": "Jar (400g)", "cost": 345.00, "stock": 1.00, "notes": "Biscoff series drinks"},
            {"sku": "SWT-005", "name": "Lotus Biscoff Biscuits (Pack 250g)", "unit": "Pack (250g)", "cost": 192.00, "stock": 1.00, "notes": "Garnish & drink topping"},
            {"sku": "SWT-006", "name": "Oreo Biscuit Crumbs 454g", "unit": "Pack (454g)", "cost": 175.00, "stock": 1.00, "notes": "Dirty Oreo & blended drinks"},
            {"sku": "SWT-007", "name": "Sprite 1.5L (Beverage Mixer)", "unit": "Bot (1.5L)", "cost": 69.50, "stock": 4.00, "notes": "Base for Fruit Soda series"},
            {"sku": "SWT-008", "name": "Fresh Lemon Fruit", "unit": "Pieces", "cost": 15.00, "stock": 42.00, "notes": "Garnish & fresh lemon drinks"},
        ]
    },
    {
        "cat_id": 6,
        "cat_name": "6. Retail Bottled Beverages & Beers",
        "icon": "🥤",
        "header_fill": "1A4870", # Deep Chiller Blue
        "items": [
            {"sku": "RET-001", "name": "Nature Spring Mineral Water 500ml", "unit": "Bot (500ml)", "cost": 12.00, "stock": 50.00, "notes": "Chiller front | SRP: ₱20.00"},
            {"sku": "RET-002", "name": "Nature Spring Mineral Water 1L", "unit": "Bot (1L)", "cost": 18.00, "stock": 20.00, "notes": "Chiller front | SRP: ₱30.00"},
            {"sku": "RET-003", "name": "Perrier Sparkling Mineral Water 330ml", "unit": "Bot (330ml)", "cost": 110.00, "stock": 12.00, "notes": "Chiller front | SRP: ₱199.00"},
            {"sku": "RET-004", "name": "Evian Natural Spring Water 500ml", "unit": "Bot (500ml)", "cost": 95.00, "stock": 12.00, "notes": "Chiller front | SRP: ₱169.00"},
            {"sku": "RET-005", "name": "Coca-Cola Regular 320ml Can", "unit": "Can (320ml)", "cost": 35.00, "stock": 24.00, "notes": "Chiller front | SRP: ₱69.00"},
            {"sku": "RET-006", "name": "Coca-Cola Zero Sugar 320ml Can", "unit": "Can (320ml)", "cost": 35.00, "stock": 24.00, "notes": "Chiller front | SRP: ₱69.00"},
            {"sku": "RET-007", "name": "Sprite Soda 320ml Can", "unit": "Can (320ml)", "cost": 35.00, "stock": 24.00, "notes": "Chiller front | SRP: ₱69.00"},
            {"sku": "RET-008", "name": "San Mig Light - Apple Bottle", "unit": "Bot (330ml)", "cost": 48.00, "stock": 12.00, "notes": "Chiller front | SRP: ₱70.00"},
            {"sku": "RET-009", "name": "San Miguel Pale Pilsen 320ml Bottle", "unit": "Bot (320ml)", "cost": 52.00, "stock": 12.00, "notes": "Chiller front | SRP: ₱80.00"},
            {"sku": "RET-010", "name": "Red Horse Stallion 330ml Bottle", "unit": "Bot (330ml)", "cost": 55.00, "stock": 15.00, "notes": "Chiller front | SRP: ₱85.00"},
        ]
    },
    {
        "cat_id": 7,
        "cat_name": "7. Pastries, Bakery & Grab-and-Go Snacks",
        "icon": "🥨",
        "header_fill": "7A3E26", # Warm Cinnamon
        "items": [
            {"sku": "PAS-001", "name": "Special Ensaymada (Pack of 6pcs)", "unit": "Pack (6s)", "cost": 45.00, "stock": 6.00, "notes": "Pastry display | SRP: ₱69.00"},
            {"sku": "PAS-002", "name": "Chocolate Crinkles Pack", "unit": "Pack", "cost": 25.00, "stock": 8.00, "notes": "Pastry display | SRP: ₱39.00"},
            {"sku": "PAS-003", "name": "Assorted Gourmet Muffins", "unit": "Pieces", "cost": 75.00, "stock": 5.00, "notes": "Pastry display | SRP: ₱129.00"},
            {"sku": "PAS-004", "name": "Sliced Pastry Cake", "unit": "Pieces", "cost": 30.00, "stock": 20.00, "notes": "Pastry display | SRP: ₱49.00"},
            {"sku": "SNK-001", "name": "Pringles Potato Crisps Original 158g", "unit": "Can (158g)", "cost": 98.00, "stock": 10.00, "notes": "Retail shelf | SRP: ₱149.00"},
            {"sku": "SNK-002", "name": "Pringles Sour Cream & Onion 158g", "unit": "Can (158g)", "cost": 98.00, "stock": 10.00, "notes": "Retail shelf | SRP: ₱149.00"},
            {"sku": "SNK-003", "name": "Pringles Cheddar Cheese 158g", "unit": "Can (158g)", "cost": 98.00, "stock": 10.00, "notes": "Retail shelf | SRP: ₱149.00"},
            {"sku": "SNK-004", "name": "Cheetos Cheddar Jalapeño 226.8g", "unit": "Pack", "cost": 98.00, "stock": 8.00, "notes": "Retail shelf | SRP: ₱149.00"},
            {"sku": "SNK-005", "name": "Butter Lover Popcorn Bag", "unit": "Bag", "cost": 85.00, "stock": 12.00, "notes": "Retail shelf | SRP: ₱149.00"},
        ]
    },
    {
        "cat_id": 8,
        "cat_name": "8. Merchandise & Collector Mugs",
        "icon": "🎁",
        "header_fill": "334A60", # Slate Indigo
        "items": [
            {"sku": "MRC-001", "name": "Mickey Mouse Ceramic Mug (with Box)", "unit": "Piece", "cost": 350.00, "stock": 20.00, "notes": "Display cabinet | SRP: ₱599.00"},
            {"sku": "MRC-002", "name": "Stitch Ceramic Mug (with Box)", "unit": "Piece", "cost": 350.00, "stock": 20.00, "notes": "Display cabinet | SRP: ₱599.00"},
            {"sku": "MRC-003", "name": "Lotso Bear Ceramic Mug (with Box)", "unit": "Piece", "cost": 350.00, "stock": 20.00, "notes": "Display cabinet | SRP: ₱599.00"},
            {"sku": "MRC-004", "name": "Nick Wilde Fox Ceramic Mug (with Box)", "unit": "Piece", "cost": 350.00, "stock": 20.00, "notes": "Display cabinet | SRP: ₱599.00"},
        ]
    },
    {
        "cat_id": 9,
        "cat_name": "9. Cups, Lids & Packaging Disposables",
        "icon": "📦",
        "header_fill": "635334", # Kraft Paper Brown
        "items": [
            {"sku": "PKG-001", "name": "Hot Paper Cups 8oz / 12oz", "unit": "Sleeve (50s)", "cost": 120.00, "stock": 5.00, "notes": "For hot coffee drinks"},
            {"sku": "PKG-002", "name": "Hot Cup Sipper Lids 80mm / 90mm", "unit": "Sleeve (50s)", "cost": 85.00, "stock": 5.00, "notes": "For hot coffee cups"},
            {"sku": "PKG-003", "name": "Cold PET Cups 16oz (Medium)", "unit": "Sleeve (50s)", "cost": 140.00, "stock": 6.00, "notes": "For iced coffee & tea 16oz"},
            {"sku": "PKG-004", "name": "Cold PET Cups 22oz (Large)", "unit": "Sleeve (50s)", "cost": 165.00, "stock": 6.00, "notes": "For iced coffee & tea 22oz"},
            {"sku": "PKG-005", "name": "Flat Lids 90mm / 95mm", "unit": "Sleeve (50s)", "cost": 75.00, "stock": 6.00, "notes": "Cold beverage flat lid"},
            {"sku": "PKG-006", "name": "Dome Lids 90mm / 95mm", "unit": "Sleeve (50s)", "cost": 85.00, "stock": 6.00, "notes": "Whipping cream / foam dome"},
            {"sku": "PKG-007", "name": "Wooden Coffee Stirrers (100s)", "unit": "Pack (100s)", "cost": 25.00, "stock": 2.00, "notes": "Eco-friendly stirrers"},
            {"sku": "PKG-008", "name": "Individually Wrapped Drinking Straws", "unit": "Pack (100s)", "cost": 45.00, "stock": 4.00, "notes": "Hygienic cold drink straws"},
            {"sku": "PKG-009", "name": "Jade Pre-Cut Table Napkins / Tissues 1000s", "unit": "Pack (1000s)", "cost": 80.00, "stock": 1.00, "notes": "Bar & dining napkins"},
            {"sku": "PKG-010", "name": "Takeout Cup Carriers (2-Cup / 4-Cup)", "unit": "Pack (50s)", "cost": 110.00, "stock": 3.00, "notes": "Kraft takeout holders"},
            {"sku": "PKG-011", "name": "Thermal POS Receipt Paper Rolls 58mm/80mm", "unit": "Rolls", "cost": 30.00, "stock": 10.00, "notes": "Cashier receipt printer"},
        ]
    },
    {
        "cat_id": 10,
        "cat_name": "10. Bar Tools, Utensils & Smallwares",
        "icon": "🥄",
        "header_fill": "4A3F50", # Dark Bronze Steel
        "items": [
            {"sku": "TOO-001", "name": "Steaming Pitcher Matte Black 600ml", "unit": "Piece", "cost": 2180.00, "stock": 1.00, "notes": "Espresso milk frothing"},
            {"sku": "TOO-002", "name": "Acrylic Measuring Jigger 10ml / 20ml", "unit": "Piece", "cost": 60.00, "stock": 2.00, "notes": "Syrup & flavor measurement"},
            {"sku": "TOO-003", "name": "Bar Spoon Glossy Stainless", "unit": "Piece", "cost": 110.00, "stock": 1.00, "notes": "Long spiral bar stirring spoon"},
            {"sku": "TOO-004", "name": "Heavy Duty Silicone Bar Mat", "unit": "Piece", "cost": 320.00, "stock": 1.00, "notes": "Counter drip station"},
            {"sku": "TOO-005", "name": "Stainless Cocktail Shaker 1L", "unit": "Piece", "cost": 250.00, "stock": 1.00, "notes": "Large iced shaken drinks"},
            {"sku": "TOO-006", "name": "Stainless Cocktail Shaker 700ml", "unit": "Piece", "cost": 210.00, "stock": 1.00, "notes": "Medium iced shaken drinks"},
            {"sku": "TOO-007", "name": "White Sauce Squeeze Bottles", "unit": "Piece", "cost": 40.00, "stock": 4.00, "notes": "Fudge & sauce squeeze dispensers"},
            {"sku": "TOO-008", "name": "Fliptop Wide Storage Containers 1.8L", "unit": "Piece", "cost": 84.00, "stock": 6.00, "notes": "Powders & topping canisters"},
            {"sku": "TOO-009", "name": "Stackable Food Keeper Square Small", "unit": "Piece", "cost": 50.00, "stock": 1.00, "notes": "Small ingredient storage"},
            {"sku": "TOO-010", "name": "Stainless Mixing Bowl", "unit": "Piece", "cost": 260.00, "stock": 1.00, "notes": "Whipping & cream mixing"},
            {"sku": "TOO-011", "name": "Ceramic Matcha Bowl / Chawan", "unit": "Piece", "cost": 90.00, "stock": 1.00, "notes": "Matcha whisking bowl"},
            {"sku": "TOO-012", "name": "Espresso Shot Glass (Double Spout)", "unit": "Piece", "cost": 22.00, "stock": 1.00, "notes": "Espresso calibration shot glass"},
            {"sku": "TOO-013", "name": "Chef Kitchen Knife Stainless", "unit": "Piece", "cost": 192.00, "stock": 1.00, "notes": "Fruit slicing & pastry cutting"},
            {"sku": "TOO-014", "name": "Heavy Duty Cutting Board", "unit": "Piece", "cost": 130.00, "stock": 1.00, "notes": "Food safe chopping board"},
            {"sku": "TOO-015", "name": "Can Opener Heavy Duty", "unit": "Piece", "cost": 129.00, "stock": 1.00, "notes": "Milk & can opener"},
            {"sku": "TOO-016", "name": "Utility Kitchen Scissors", "unit": "Piece", "cost": 119.00, "stock": 1.00, "notes": "Pack opening & prep"},
            {"sku": "TOO-017", "name": "Non-Slip Serving Trays", "unit": "Pack (2s)", "cost": 156.00, "stock": 2.00, "notes": "Dine-in tray service"},
            {"sku": "TOO-018", "name": "Condiment Organizer 4-Compartment", "unit": "Piece", "cost": 420.00, "stock": 1.00, "notes": "Straw & napkin counter caddy"},
            {"sku": "TOO-019", "name": "Cup & Lid Dispenser Holder", "unit": "Piece", "cost": 498.00, "stock": 1.00, "notes": "Countertop cup carousel"},
        ]
    },
    {
        "cat_id": 11,
        "cat_name": "11. Equipment & Store Fixtures",
        "icon": "⚙️",
        "header_fill": "26384A", # Modern Slate Charcoal
        "items": [
            {"sku": "EQP-001", "name": "Electric Hand Mixer (High Power)", "unit": "Piece", "cost": 1510.00, "stock": 1.00, "notes": "Cheese foam / whipping"},
            {"sku": "EQP-002", "name": "AWD DO3 Handheld Drink Frother", "unit": "Piece", "cost": 310.00, "stock": 1.00, "notes": "Single cup frothing"},
            {"sku": "EQP-003", "name": "Commercial Insulated Ice Box Chest", "unit": "Piece", "cost": 784.00, "stock": 1.00, "notes": "Bar ice storage container"},
            {"sku": "EQP-004", "name": "Heavy Duty Steel Cash Box with Lock", "unit": "Piece", "cost": 530.00, "stock": 1.00, "notes": "POS register petty cash box"},
            {"sku": "EQP-005", "name": "Akari Heavy Duty Power Adaptor", "unit": "Piece", "cost": 69.00, "stock": 1.00, "notes": "Appliance surge plug"},
            {"sku": "EQP-006", "name": "Universal Surge Protected Power Strip", "unit": "Piece", "cost": 298.00, "stock": 1.00, "notes": "POS & equipment power strip"},
            {"sku": "EQP-007", "name": "Barista Recipe Manual & Store Record Book", "unit": "Piece", "cost": 72.00, "stock": 1.00, "notes": "Official store logbook"},
            {"sku": "EQP-008", "name": "Commercial Foot-Pedal Trash Can 30L", "unit": "Piece", "cost": 850.00, "stock": 1.00, "notes": "Sanitary waste disposal"},
        ]
    },
    {
        "cat_id": 12,
        "cat_name": "12. Cleaning, Hygiene & Maintenance",
        "icon": "🧼",
        "header_fill": "125252", # Fresh Aqua Marine
        "items": [
            {"sku": "CLN-001", "name": "Dr. Coffee Espresso Descaling Powder", "unit": "Bot", "cost": 1000.00, "stock": 1.00, "notes": "Espresso boiler descaling"},
            {"sku": "CLN-002", "name": "Dr. Coffee Milk Line Deep Clean Solution", "unit": "Bot", "cost": 2800.00, "stock": 1.00, "notes": "Steam wand & pipe sanitizer"},
            {"sku": "CLN-003", "name": "Microfiber Bar Towels / Cleaning Cloths", "unit": "Piece", "cost": 115.00, "stock": 2.00, "notes": "Bar wipe & steam wand towel"},
            {"sku": "CLN-004", "name": "Heavy Duty Black Garbage Bags 26x32 (50s)", "unit": "Pack (50s)", "cost": 300.00, "stock": 1.00, "notes": "Waste bin liners"},
            {"sku": "CLN-005", "name": "Heavy Duty Packaging & Masking Tape", "unit": "Piece", "cost": 159.00, "stock": 1.00, "notes": "Box sealing & date labels"},
            {"sku": "CLN-006", "name": "Whiteboard Marker & Shelf Label Pens", "unit": "Piece", "cost": 64.00, "stock": 1.00, "notes": "Store menu board & dating"},
            {"sku": "CLN-007", "name": "Store Promotional Balloons & Hand Pump", "unit": "Set", "cost": 104.00, "stock": 1.00, "notes": "Marketing & events kit"},
        ]
    }
]

# --- OPENPYXL STYLES SETUP ---
font_family = "Segoe UI"
title_font = Font(name=font_family, size=16, bold=True, color="2B1704")
subtitle_font = Font(name=font_family, size=10, italic=False, color="5A4A3A")

meta_label_font = Font(name=font_family, size=9, bold=True, color="555555")
meta_val_font = Font(name=font_family, size=10, bold=True, color="1F1F1F")

kpi_num_font = Font(name=font_family, size=14, bold=True, color="1F1F1F")
kpi_lbl_font = Font(name=font_family, size=8, bold=True, color="666666")

header_col_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
input_header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
cat_header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
subtotal_font = Font(name=font_family, size=10, bold=True, color="2B1704")

data_font = Font(name=font_family, size=10, color="222222")
data_bold_font = Font(name=font_family, size=10, bold=True, color="111111")
notes_font = Font(name=font_family, size=9, italic=True, color="666666")

# Borders
border_light_gray = Side(border_style="thin", color="E0DCD5")
cell_border = Border(left=border_light_gray, right=border_light_gray, top=border_light_gray, bottom=border_light_gray)

thick_caramel = Side(border_style="medium", color="C68B59")
input_border = Border(left=thick_caramel, right=thick_caramel, top=border_light_gray, bottom=border_light_gray)

subtotal_top = Side(border_style="thin", color="4A2E18")
subtotal_bot = Side(border_style="double", color="4A2E18")
subtotal_border = Border(left=border_light_gray, right=border_light_gray, top=subtotal_top, bottom=subtotal_bot)

kpi_border = Border(left=Side(border_style="thin", color="B0A695"),
                    right=Side(border_style="thin", color="B0A695"),
                    top=Side(border_style="thin", color="B0A695"),
                    bottom=Side(border_style="thin", color="B0A695"))

# Fills
main_header_fill = PatternFill(start_color="2B1704", end_color="2B1704", fill_type="solid") # Dark Roast Espresso
input_col_header_fill = PatternFill(start_color="D48806", end_color="D48806", fill_type="solid") # Warm Caramel Callout
zebra_light_fill = PatternFill(start_color="FDFBF7", end_color="FDFBF7", fill_type="solid") # Light Cream
zebra_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
input_cell_fill = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid") # Pale Canary Yellow / Cream Input
subtotal_fill = PatternFill(start_color="F5EBE6", end_color="F5EBE6", fill_type="solid") # Soft Mocha
meta_card_fill = PatternFill(start_color="F9F6F0", end_color="F9F6F0", fill_type="solid")

kpi_blue_fill = PatternFill(start_color="EDF4FA", end_color="EDF4FA", fill_type="solid")
kpi_green_fill = PatternFill(start_color="EAF6ED", end_color="EAF6ED", fill_type="solid")
kpi_amber_fill = PatternFill(start_color="FEF8EA", end_color="FEF8EA", fill_type="solid")
kpi_red_fill = PatternFill(start_color="FDF0EE", end_color="FDF0EE", fill_type="solid")


def build_inventory_sheet(ws, title_text, included_categories, sheet_type="master"):
    ws.views.sheetView[0].showGridLines = True

    # 1. Store Header & Title
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "☕ ESPRESSO YOURSELF & TEA HOUSE — PHYSICAL INVENTORY AUDIT"
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Branch: Cebu City (#30)  |  Audit Section: {title_text}  |  Generated for Today's Stock Count"
    sub_cell.font = subtitle_font
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # 2. Metadata Block (Date, Barista, Shift, Auditor)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 22

    metadata_fields = [
        ("A4", "B4", "AUDIT DATE:", "C4", datetime.now().strftime("%B %d, %Y")),
        ("A5", "B5", "STORE SHIFT:", "C5", "All-Day / Closing Audit"),
        ("E4", "F4", "COUNTED BY (BARISTA):", "G4", "________________________"),
        ("E5", "F5", "VERIFIED BY (SUPERVISOR):", "G5", "________________________"),
        ("I4", "J4", "STATUS:", "K4", "READY FOR COUNT"),
        ("I5", "J5", "DISCREPANCY ACTION:", "K5", "Requires Manager Sign-off"),
    ]

    for lbl_start, lbl_end, lbl_text, val_cell, val_text in metadata_fields:
        ws.merge_cells(f"{lbl_start}:{lbl_end}")
        c_lbl = ws[lbl_start]
        c_lbl.value = lbl_text
        c_lbl.font = meta_label_font
        c_lbl.alignment = Alignment(horizontal="right", vertical="center")
        c_lbl.fill = meta_card_fill

        c_val = ws[val_cell]
        c_val.value = val_text
        c_val.font = meta_val_font
        c_val.alignment = Alignment(horizontal="left", vertical="center")

    # 3. Column Headers
    headers = [
        ("A7", "SKU / CODE", 14),
        ("B7", "ITEM DESCRIPTION & SPECS", 38),
        ("C7", "UOM", 14),
        ("D7", "UNIT COST (₱)", 15),
        ("E7", "SYSTEM STOCK", 15),
        ("F7", "★ PHYSICAL COUNT ★", 21), # Designated user input column
        ("G7", "VARIANCE (QTY)", 16),
        ("H7", "VARIANCE VALUE (₱)", 18),
        ("I7", "TOTAL COUNT VALUE (₱)", 20),
        ("J7", "AUDIT STATUS", 16),
        ("K7", "BARISTA REMARKS / EXPIRY", 32)
    ]

    ws.row_dimensions[7].height = 28
    for col_cell, col_name, width in headers:
        cell = ws[col_cell]
        cell.value = col_name
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        
        # Highlight input column
        if "PHYSICAL COUNT" in col_name:
            cell.font = input_header_font
            cell.fill = input_col_header_fill
        else:
            cell.font = header_col_font
            cell.fill = main_header_fill

        col_letter = col_cell[0]
        ws.column_dimensions[col_letter].width = width

    # 4. Populate Categories and Items
    current_row = 8
    category_summary_rows = []
    all_item_rows = []

    for cat in included_categories:
        cat_name = cat["cat_name"]
        cat_color = cat["header_fill"]
        icon = cat.get("icon", "📂")
        cat_fill = PatternFill(start_color=cat_color, end_color=cat_color, fill_type="solid")
        
        # Category Banner Row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
        cat_cell = ws.cell(row=current_row, column=1, value=f"  {icon}  {cat_name.upper()}")
        cat_cell.font = cat_header_font
        cat_cell.fill = cat_fill
        cat_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 24
        current_row += 1

        cat_start_row = current_row

        for idx, item in enumerate(cat["items"]):
            r = current_row
            all_item_rows.append(r)
            ws.row_dimensions[r].height = 21

            sku_cell = ws.cell(row=r, column=1, value=item["sku"])
            name_cell = ws.cell(row=r, column=2, value=item["name"])
            uom_cell = ws.cell(row=r, column=3, value=item["unit"])
            cost_cell = ws.cell(row=r, column=4, value=item["cost"])
            sys_cell = ws.cell(row=r, column=5, value=item["stock"])
            
            # Column 6 (F): User Input - Blank initially for Baristas to input
            phys_cell = ws.cell(row=r, column=6)
            phys_cell.value = None

            # Column 7 (G): Variance Qty Formula: Physical - System Stock
            var_qty_cell = ws.cell(row=r, column=7)
            var_qty_cell.value = f'=IF(ISBLANK(F{r}), "", F{r} - E{r})'

            # Column 8 (H): Variance Value Formula: Variance Qty * Unit Cost
            var_val_cell = ws.cell(row=r, column=8)
            var_val_cell.value = f'=IF(ISBLANK(F{r}), "", G{r} * D{r})'

            # Column 9 (I): Total Count Value Formula: Physical Count * Unit Cost
            tot_val_cell = ws.cell(row=r, column=9)
            tot_val_cell.value = f'=IF(ISBLANK(F{r}), "", F{r} * D{r})'

            # Column 10 (J): Audit Status Badge Formula
            status_cell = ws.cell(row=r, column=10)
            status_cell.value = f'=IF(ISBLANK(F{r}), "⏳ Pending", IF(G{r}=0, "✅ Match", IF(G{r}<0, "⚠️ Shortage", "🔺 Surplus")))'

            # Column 11 (K): Remarks
            rem_cell = ws.cell(row=r, column=11, value=item.get("notes", ""))

            # Format numbers and alignments
            sku_cell.alignment = Alignment(horizontal="center", vertical="center")
            name_cell.alignment = Alignment(horizontal="left", vertical="center")
            uom_cell.alignment = Alignment(horizontal="center", vertical="center")
            cost_cell.alignment = Alignment(horizontal="right", vertical="center")
            sys_cell.alignment = Alignment(horizontal="right", vertical="center")
            phys_cell.alignment = Alignment(horizontal="center", vertical="center")
            var_qty_cell.alignment = Alignment(horizontal="right", vertical="center")
            var_val_cell.alignment = Alignment(horizontal="right", vertical="center")
            tot_val_cell.alignment = Alignment(horizontal="right", vertical="center")
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            rem_cell.alignment = Alignment(horizontal="left", vertical="center")

            # Number Formats
            cost_cell.number_format = '₱#,##0.00'
            sys_cell.number_format = '#,##0.00'
            phys_cell.number_format = '#,##0.00'
            var_qty_cell.number_format = '+#,##0.00;-#,##0.00;0.00'
            var_val_cell.number_format = '₱+#,##0.00;₱-#,##0.00;₱0.00'
            tot_val_cell.number_format = '₱#,##0.00'

            # Fonts
            sku_cell.font = Font(name=font_family, size=9, color="666666")
            name_cell.font = data_bold_font
            uom_cell.font = data_font
            cost_cell.font = data_font
            sys_cell.font = data_font
            phys_cell.font = Font(name=font_family, size=11, bold=True, color="9C5B00")
            var_qty_cell.font = data_bold_font
            var_val_cell.font = data_bold_font
            tot_val_cell.font = data_bold_font
            status_cell.font = data_bold_font
            rem_cell.font = notes_font

            # Fills & Borders
            row_fill = zebra_light_fill if idx % 2 == 1 else zebra_white_fill
            for col_idx in [1, 2, 3, 4, 5, 7, 8, 9, 10, 11]:
                c = ws.cell(row=r, column=col_idx)
                c.fill = row_fill
                c.border = cell_border

            # User input column gets distinct highlighting
            phys_cell.fill = input_cell_fill
            phys_cell.border = input_border

            current_row += 1

        cat_end_row = current_row - 1

        # Category Subtotal Row
        ws.row_dimensions[current_row].height = 22
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        sub_lbl = ws.cell(row=current_row, column=1, value=f"Subtotal — {cat_name.split('. ')[-1]}:")
        sub_lbl.font = subtotal_font
        sub_lbl.alignment = Alignment(horizontal="right", vertical="center")

        sub_cost = ws.cell(row=current_row, column=4, value="")
        sub_sys = ws.cell(row=current_row, column=5, value=f"=SUM(E{cat_start_row}:E{cat_end_row})")
        sub_phys = ws.cell(row=current_row, column=6, value=f"=SUM(F{cat_start_row}:F{cat_end_row})")
        sub_var_qty = ws.cell(row=current_row, column=7, value=f"=SUM(G{cat_start_row}:G{cat_end_row})")
        sub_var_val = ws.cell(row=current_row, column=8, value=f"=SUM(H{cat_start_row}:H{cat_end_row})")
        sub_tot_val = ws.cell(row=current_row, column=9, value=f"=SUM(I{cat_start_row}:I{cat_end_row})")
        sub_status = ws.cell(row=current_row, column=10, value="")
        sub_rem = ws.cell(row=current_row, column=11, value="")

        sub_sys.number_format = '#,##0.00'
        sub_phys.number_format = '#,##0.00'
        sub_var_qty.number_format = '+#,##0.00;-#,##0.00;0.00'
        sub_var_val.number_format = '₱+#,##0.00;₱-#,##0.00;₱0.00'
        sub_tot_val.number_format = '₱#,##0.00'

        for c in [sub_sys, sub_phys, sub_var_qty, sub_var_val, sub_tot_val]:
            c.font = subtotal_font
            c.alignment = Alignment(horizontal="right", vertical="center")

        for col_idx in range(1, 12):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = subtotal_fill
            c.border = subtotal_border

        category_summary_rows.append(current_row)
        current_row += 2

    # 5. Grand Total Row at Bottom
    grand_row = current_row
    ws.row_dimensions[grand_row].height = 26
    ws.merge_cells(start_row=grand_row, start_column=1, end_row=grand_row, end_column=3)
    grand_lbl = ws.cell(row=grand_row, column=1, value="🏆 GRAND TOTAL (AUDIT SUMMARY):")
    grand_lbl.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    grand_lbl.alignment = Alignment(horizontal="right", vertical="center")

    # Use SUM(cell1, cell2, ...) to safely handle any blanks/strings without #VALUE! error
    if category_summary_rows:
        sys_args = ", ".join([f"E{r}" for r in category_summary_rows])
        phys_args = ", ".join([f"F{r}" for r in category_summary_rows])
        var_qty_args = ", ".join([f"G{r}" for r in category_summary_rows])
        var_val_args = ", ".join([f"H{r}" for r in category_summary_rows])
        tot_val_args = ", ".join([f"I{r}" for r in category_summary_rows])

        g_sys = ws.cell(row=grand_row, column=5, value=f"=SUM({sys_args})")
        g_phys = ws.cell(row=grand_row, column=6, value=f"=SUM({phys_args})")
        g_var_qty = ws.cell(row=grand_row, column=7, value=f"=SUM({var_qty_args})")
        g_var_val = ws.cell(row=grand_row, column=8, value=f"=SUM({var_val_args})")
        g_tot_val = ws.cell(row=grand_row, column=9, value=f"=SUM({tot_val_args})")

        g_sys.number_format = '#,##0.00'
        g_phys.number_format = '#,##0.00'
        g_var_qty.number_format = '+#,##0.00;-#,##0.00;0.00'
        g_var_val.number_format = '₱+#,##0.00;₱-#,##0.00;₱0.00'
        g_tot_val.number_format = '₱#,##0.00'

        for c in [g_sys, g_phys, g_var_qty, g_var_val, g_tot_val]:
            c.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 12):
        c = ws.cell(row=grand_row, column=col_idx)
        c.fill = main_header_fill
        c.border = cell_border

    # Conditional Formatting for Variance (Column G & H) and Status (Column J)
    if all_item_rows:
        first_r = all_item_rows[0]
        last_r = all_item_rows[-1]

        # Red fill for Shortage (< 0)
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        red_font = Font(name=font_family, size=10, bold=True, color="C00000")
        ws.conditional_formatting.add(
            f"G{first_r}:H{last_r}",
            CellIsRule(operator="lessThan", formula=["0"], stopIfTrue=True, fill=red_fill, font=red_font)
        )

        # Soft green fill for Match / Surplus (>= 0)
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        green_font = Font(name=font_family, size=10, bold=True, color="375623")
        ws.conditional_formatting.add(
            f"G{first_r}:H{last_r}",
            CellIsRule(operator="greaterThanOrEqual", formula=["0"], stopIfTrue=True, fill=green_fill, font=green_font)
        )

    # Sign-off verification footer
    sign_row = grand_row + 3
    ws.cell(row=sign_row, column=2, value="Conducted By (Barista Signature): ____________________").font = data_bold_font
    ws.cell(row=sign_row, column=6, value="Audited & Accepted By (Manager Signature): ____________________").font = data_bold_font
    ws.cell(row=sign_row+1, column=2, value="Time Completed: ____________________").font = data_font
    ws.cell(row=sign_row+1, column=6, value="Date Approved: ____________________").font = data_font


def generate_workbook():
    print("Creating Excel Workbook...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Tab 1: Master Inventory Count (All 12 Categories)
    ws_master = wb.create_sheet(title="📋 Master Inventory Count")
    build_inventory_sheet(ws_master, "Full Store Master Audit", CATEGORIES_DATA, sheet_type="master")

    # Tab 2: Bar & Kitchen (Raw Ingredients & Packaging)
    bar_categories = [cat for cat in CATEGORIES_DATA if cat["cat_id"] in [1, 2, 3, 4, 5, 9]]
    ws_bar = wb.create_sheet(title="☕ Bar & Kitchen (Ingredients)")
    build_inventory_sheet(ws_bar, "Bar Consumables & Packaging", bar_categories, sheet_type="bar")

    # Tab 3: Retail, Drinks & Snacks
    retail_categories = [cat for cat in CATEGORIES_DATA if cat["cat_id"] in [6, 7, 8]]
    ws_retail = wb.create_sheet(title="🥤 Retail, Drinks & Merch")
    build_inventory_sheet(ws_retail, "Chiller, Pastry & Retail Count", retail_categories, sheet_type="retail")

    # Tab 4: Equipment & Smallwares
    equipment_categories = [cat for cat in CATEGORIES_DATA if cat["cat_id"] in [10, 11, 12]]
    ws_eqp = wb.create_sheet(title="🛠️ Equipment & Supplies")
    build_inventory_sheet(ws_eqp, "Bar Tools, Equipment & Cleaning", equipment_categories, sheet_type="equipment")

    # Tab 5: Quick Walkaround Printable Sheet
    ws_print = wb.create_sheet(title="🖨️ Printable Walkaround Sheet")
    build_inventory_sheet(ws_print, "Clipboard Quick Count (Pen & Paper)", CATEGORIES_DATA, sheet_type="printable")
    ws_print.page_setup.orientation = ws_print.ORIENTATION_LANDSCAPE
    ws_print.page_setup.paperSize = ws_print.PAPERSIZE_LETTER
    ws_print.page_setup.fitToPage = True
    ws_print.page_setup.fitToWidth = 1
    ws_print.page_setup.fitToHeight = 0

    # Save to recordsExcel
    wb.save(excel_output_path)
    print(f"[SUCCESS] Saved template to: {excel_output_path}")

    # Also save copy to workspace root
    wb.save(workspace_copy_path)
    print(f"[SUCCESS] Saved copy to: {workspace_copy_path}")


if __name__ == "__main__":
    generate_workbook()
