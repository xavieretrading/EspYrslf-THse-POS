import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# File paths
rec_dir = r"c:\Users\Philippines Freight\MainSystems\POS\recordsExcel"
workspace_dir = r"c:\Users\Philippines Freight\MainSystems\POS"
excel_path = os.path.join(rec_dir, "Espresso_Yourself_Inventory_Audit_with_Beginning_Stocks.xlsx")
workspace_copy_path = os.path.join(workspace_dir, "Espresso_Yourself_Inventory_Audit_with_Beginning_Stocks.xlsx")

# Load existing workbook
wb = openpyxl.load_workbook(excel_path)

# Palette
CLR_PRIMARY_DARK = "2C1810"    # Deep Espresso
CLR_PRIMARY_MED = "4A2E1B"     # Roast Brown
CLR_ACCENT_GOLD = "C8963E"     # Warm Caramel Gold

FONT_TITLE = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Segoe UI", size=10, italic=True, color="E0D4C3")
FONT_HDR = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Segoe UI", size=10, color="1F1F1F")
FONT_DATA_BOLD = Font(name="Segoe UI", size=10, bold=True, color="1F1F1F")
FONT_NOTE = Font(name="Segoe UI", size=9, italic=True, color="555555")

FILL_BEST = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")       # Light Green
FONT_BEST = Font(name="Segoe UI", size=10, bold=True, color="0F5132")

FILL_MOD = PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")        # Light Blue
FONT_MOD = Font(name="Segoe UI", size=10, bold=True, color="084298")

FILL_ZERO = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")       # Light Gray
FONT_ZERO = Font(name="Segoe UI", size=10, color="6C757D")

THIN_SIDE = Side(border_style="thin", color="D3D3D3")
BORDER_DATA = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_HEADER = Border(left=Side(style="thin", color="555555"), right=Side(style="thin", color="555555"), top=Side(style="medium", color="2C1810"), bottom=Side(style="medium", color="2C1810"))
BORDER_TOTAL = Border(top=Side(style="thin", color="2C1810"), bottom=Side(style="double", color="2C1810"))

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Load beverage data
with open("scratch/beverage_menu_sales.json", "r", encoding="utf-8") as f:
    data = json.load(f)

grouped = data["grouped"]
cats_order = [
    ("Hot Coffee", "☕", "3D2314"),
    ("Iced & Blended", "🧊", "1A4870"),
    ("Soda Based", "🫧", "6B3E26"),
    ("Fruit Teas", "🍓", "7A3E26"),
    ("Matcha Series", "🍵", "2D5A27"),
    ("Frappe", "🥤", "7A5210"),
    ("TEA", "🫖", "1B4D3E")
]

# Remove existing sheet if re-running
if "☕ Beverage Sales & Variance" in wb.sheetnames:
    wb.remove(wb["☕ Beverage Sales & Variance"])

ws = wb.create_sheet(title="☕ Beverage Sales & Variance", index=2)
ws.views.sheetView[0].showGridLines = True

# Title Header
ws.merge_cells("A1:I1")
ws["A1"] = "ESPRESSO YOURSELF & TEA HOUSE — PREPARED BEVERAGE INVENTORY & SALES AUDIT"
ws["A1"].font = FONT_TITLE
ws["A1"].fill = PatternFill(start_color=CLR_PRIMARY_DARK, end_color=CLR_PRIMARY_DARK, fill_type="solid")
ws["A1"].alignment = ALIGN_CENTER
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:I2")
ws["A2"] = "Branch #30 (Cebu City)  |  Beginning System Baseline vs Current Stock Now = Sales Units Sold  |  Audit Date: September 22, 2026"
ws["A2"].font = FONT_SUBTITLE
ws["A2"].fill = PatternFill(start_color=CLR_PRIMARY_MED, end_color=CLR_PRIMARY_MED, fill_type="solid")
ws["A2"].alignment = ALIGN_CENTER
ws.row_dimensions[2].height = 22

headers = [
    ("A", "Product ID", 12, ALIGN_CENTER),
    ("B", "Beverage / Coffee Name", 36, ALIGN_LEFT),
    ("C", "Category", 18, ALIGN_LEFT),
    ("D", "Unit", 10, ALIGN_CENTER),
    ("E", "Price (₱)", 12, ALIGN_RIGHT),
    ("F", "Beginning Stock\n(Baseline)", 16, ALIGN_RIGHT),
    ("G", "★ Stock Now ★\n(Current POS)", 16, ALIGN_RIGHT),
    ("H", "Sales (Variance)\n(Units Sold)", 16, ALIGN_RIGHT),
    ("I", "Total Sales\nRevenue (₱)", 18, ALIGN_RIGHT),
]

ws.row_dimensions[4].height = 28
for col_let, h_text, width, align in headers:
    cell = ws[f"{col_let}4"]
    cell.value = h_text
    cell.font = FONT_HDR
    cell.fill = PatternFill(start_color="3D2314", end_color="3D2314", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER_HEADER
    ws.column_dimensions[col_let].width = width

current_row = 5
subtotal_rows = []

for cat_name, icon, hdr_clr in cats_order:
    prods = grouped.get(cat_name, [])
    if not prods:
        continue
        
    # Sort products so highest sales are at the top
    prods.sort(key=lambda p: (100 if p["stock"] <= 100 else 9999) - p["stock"], reverse=True)
    
    # Category Header Banner
    ws.merge_cells(f"A{current_row}:I{current_row}")
    sec_cell = ws[f"A{current_row}"]
    sec_cell.value = f"{icon}  CATEGORY: {cat_name.upper()} ({len(prods)} Beverages)"
    sec_cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    sec_cell.fill = PatternFill(start_color=hdr_clr, end_color=hdr_clr, fill_type="solid")
    sec_cell.alignment = ALIGN_LEFT
    ws.row_dimensions[current_row].height = 24
    
    start_sec_row = current_row + 1
    current_row += 1
    
    for p in prods:
        ws.row_dimensions[current_row].height = 20
        cur_stock = p["stock"]
        
        # Determine Beginning Baseline:
        # If Dilmah tea (<=100): 100
        # If regular drink (<=9999): 9999
        # If > 9999: 10000
        if cur_stock <= 100:
            beg_val = 100
        elif cur_stock > 9999:
            beg_val = 10000
        else:
            beg_val = 9999
            
        # Col A: ID
        ws.cell(current_row, 1, f"DRK-{p['id']}").alignment = ALIGN_CENTER
        ws.cell(current_row, 1).font = Font(name="Segoe UI", size=9, bold=True, color="555555")
        
        # Col B: Name
        ws.cell(current_row, 2, p["name"]).alignment = ALIGN_LEFT
        ws.cell(current_row, 2).font = FONT_DATA_BOLD
        
        # Col C: Category
        ws.cell(current_row, 3, cat_name).alignment = ALIGN_LEFT
        ws.cell(current_row, 3).font = FONT_DATA
        
        # Col D: Unit
        ws.cell(current_row, 4, p.get("unit", "cups")).alignment = ALIGN_CENTER
        ws.cell(current_row, 4).font = FONT_DATA
        
        # Col E: Price
        price_cell = ws.cell(current_row, 5, p["price"])
        price_cell.alignment = ALIGN_RIGHT
        price_cell.font = FONT_DATA
        price_cell.number_format = '₱#,##0.00'
        
        # Col F: Beginning Stock (Baseline)
        beg_cell = ws.cell(current_row, 6, beg_val)
        beg_cell.alignment = ALIGN_RIGHT
        beg_cell.font = FONT_DATA
        beg_cell.number_format = '#,##0'
        
        # Col G: Current Stock Now
        now_cell = ws.cell(current_row, 7, cur_stock)
        now_cell.alignment = ALIGN_RIGHT
        now_cell.font = FONT_DATA_BOLD
        now_cell.number_format = '#,##0'
        
        # Col H: Sales Variance = Beginning - Stock Now
        # (e.g. 9999 - 9994 = 5)
        sales_cell = ws.cell(current_row, 8, f"=F{current_row}-G{current_row}")
        sales_cell.alignment = ALIGN_RIGHT
        sales_cell.font = FONT_DATA_BOLD
        sales_cell.number_format = '#,##0'
        
        # Highlight sales
        sales_calc = beg_val - cur_stock
        if sales_calc >= 20:
            sales_cell.fill = FILL_BEST
            sales_cell.font = FONT_BEST
        elif sales_calc > 0:
            sales_cell.fill = FILL_MOD
            sales_cell.font = FONT_MOD
        else:
            sales_cell.fill = FILL_ZERO
            sales_cell.font = FONT_ZERO
            
        # Col I: Total Sales Revenue = Sales * Price
        rev_cell = ws.cell(current_row, 9, f"=H{current_row}*E{current_row}")
        rev_cell.alignment = ALIGN_RIGHT
        rev_cell.font = FONT_DATA_BOLD
        rev_cell.number_format = '₱#,##0.00'
        
        row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid") if current_row % 2 == 0 else PatternFill(start_color="FAF8F5", end_color="FAF8F5", fill_type="solid")
        for c in range(1, 10):
            cell = ws.cell(current_row, c)
            cell.border = BORDER_DATA
            if c not in [8]: # preserve sales highlight
                cell.fill = row_fill
                
        current_row += 1
        
    # Subtotal Row
    ws.row_dimensions[current_row].height = 22
    ws.merge_cells(f"A{current_row}:G{current_row}")
    sub_title = ws.cell(current_row, 1, f"Subtotal — {cat_name} Sales:")
    sub_title.alignment = Alignment(horizontal="right", vertical="center")
    sub_title.font = Font(name="Segoe UI", size=9, bold=True, color=CLR_PRIMARY_MED)
    
    # Subtotal Sales Units Sold
    sub_sales = ws.cell(current_row, 8, f"=SUM(H{start_sec_row}:H{current_row-1})")
    sub_sales.alignment = ALIGN_RIGHT
    sub_sales.font = Font(name="Segoe UI", size=10, bold=True, color="084298")
    sub_sales.number_format = '#,##0'
    
    # Subtotal Revenue
    sub_rev = ws.cell(current_row, 9, f"=SUM(I{start_sec_row}:I{current_row-1})")
    sub_rev.alignment = ALIGN_RIGHT
    sub_rev.font = Font(name="Segoe UI", size=10, bold=True, color="0F5132")
    sub_rev.number_format = '₱#,##0.00'
    
    subtotal_rows.append(current_row)
    
    for c in range(1, 10):
        ws.cell(current_row, c).border = Border(top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
        ws.cell(current_row, c).fill = PatternFill(start_color="F2EDE4", end_color="F2EDE4", fill_type="solid")
        
    current_row += 1

# Grand Total Row
ws.row_dimensions[current_row].height = 28
ws.merge_cells(f"A{current_row}:G{current_row}")
gt_cell = ws.cell(current_row, 1, "🏆 GRAND TOTAL — ALL BEVERAGE SALES:")
gt_cell.font = Font(name="Segoe UI", size=11, bold=True, color=CLR_PRIMARY_DARK)
gt_cell.alignment = Alignment(horizontal="right", vertical="center")

# Sum of all category subtotal sales
sales_sum_formula = "=" + "+".join([f"H{r}" for r in subtotal_rows])
gt_sales = ws.cell(current_row, 8, sales_sum_formula)
gt_sales.font = Font(name="Segoe UI", size=12, bold=True, color="084298")
gt_sales.alignment = ALIGN_RIGHT
gt_sales.number_format = '#,##0'

# Sum of all category subtotal revenues
rev_sum_formula = "=" + "+".join([f"I{r}" for r in subtotal_rows])
gt_rev = ws.cell(current_row, 9, rev_sum_formula)
gt_rev.font = Font(name="Segoe UI", size=12, bold=True, color="0F5132")
gt_rev.alignment = ALIGN_RIGHT
gt_rev.number_format = '₱#,##0.00'

for c in range(1, 10):
    ws.cell(current_row, c).border = BORDER_TOTAL
    ws.cell(current_row, c).fill = PatternFill(start_color="E6DFD5", end_color="E6DFD5", fill_type="solid")

ws.freeze_panes = "A5"

# Save updated workbook to dedicated new file and attempt updates
new_full_path = os.path.join(rec_dir, "Espresso_Yourself_Full_Inventory_and_Coffee_Sales_Audit.xlsx")
new_full_workspace = os.path.join(workspace_dir, "Espresso_Yourself_Full_Inventory_and_Coffee_Sales_Audit.xlsx")

wb.save(new_full_path)
wb.save(new_full_workspace)
print(f"Successfully generated: {new_full_path}")
print(f"Successfully mirrored: {new_full_workspace}")

# Attempt to update the other files if not locked
for p in [excel_path, workspace_copy_path, os.path.join(rec_dir, "Espresso_Yourself_Physical_vs_POS_Inventory_Audit.xlsx")]:
    try:
        wb.save(p)
        print(f"Also updated: {p}")
    except PermissionError:
        print(f"File {p} currently locked by Excel.")

