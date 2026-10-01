import sys, os
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\Philippines Freight\MainSystems\POS\recordsExcel\InventoryEspresso\Espresso-Beginning Stocks (5).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n==========================================")
    print(f"SHEET: {sname} (Rows: {ws.max_row})")
    print(f"==========================================")
    for r in range(1, ws.max_row + 1):
        vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        if any(vals):
            print(f"Row {r:2d}: {' | '.join(vals)}")
