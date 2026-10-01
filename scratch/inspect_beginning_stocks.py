import sys, os
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\Philippines Freight\MainSystems\POS\recordsExcel\InventoryEspresso\Espresso-Beginning Stocks (5).xlsx"

wb = openpyxl.load_workbook(file_path, data_only=True)
print(f"Sheet names in {os.path.basename(file_path)}:")
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=== Sheet: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ===")
    for r in range(1, min(40, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or '') for c in range(1, min(20, ws.max_column + 1))]
        if any(v.strip() for v in row_vals):
            print(f"Row {r:2d}: {' | '.join(row_vals[:12])}")
