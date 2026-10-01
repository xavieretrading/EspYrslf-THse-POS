import openpyxl

wb = openpyxl.load_workbook('Espresso_Branch_Inventory_Template.xlsx')
with open('scratch/excel_template_summary.txt', 'w', encoding='utf-8') as out:
    for name in wb.sheetnames:
        sheet = wb[name]
        out.write(f"=== Sheet: {name} (Rows: {sheet.max_row}, Cols: {sheet.max_column}) ===\n")
        for r in range(1, min(60, sheet.max_row + 1)):
            row_vals = [str(sheet.cell(r, c).value or '') for c in range(1, min(15, sheet.max_column + 1))]
            if any(row_vals):
                out.write(f"Row {r:2d}: {' | '.join(row_vals)}\n")
print("Wrote excel_template_summary.txt")
