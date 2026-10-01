import sys, os
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\Philippines Freight\MainSystems\POS\recordsExcel\InventoryEspresso\Espresso-Beginning Stocks (5).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print("Searching for cups/packaging across sheets:")
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '').lower()
            if any(k in val for k in ['cup', 'plastic', 'paper', '16oz', '8oz']):
                row_vals = [str(ws.cell(r, col).value or '') for col in range(1, ws.max_column + 1)]
                print(f"[{sname}] Row {r}: {' | '.join(row_vals)}")
